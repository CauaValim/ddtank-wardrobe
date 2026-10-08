"""
Gera src/lib/eventTemplate/codesManifest.json a partir do modelo de solicitação de códigos.

Uso:
    python3 scripts/code-template/build_codes_manifest.py <modelo-codigos.xlsx> <saida.json>

Cada aba do modelo vira um layout: KickSub e Lives têm um bloco; as abas de torneio têm um bloco
por premiação, lado a lado (8 colunas cada). As justificativas ficam como estão no modelo.
"""
import hashlib, json, re, sys
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

SRC, OUT = sys.argv[1], sys.argv[2]
VERSION = "codigos-v1"
FILE_NAME = "modelo-codigos-v1.xlsx"
wb = openpyxl.load_workbook(SRC)


def col(c, dc):
    return L(CI(c) + dc)


def text(ws, ref):
    v = ws[ref].value
    return "" if v is None else str(v)


def field(key, label, cell, inp, fmt="{value}", rich="plain", default=None):
    f = {"key": key, "label": label, "cell": cell, "input": inp, "format": fmt, "rich": rich}
    if default:
        f["default"] = default
    return f


def activity_name(ws, ref):
    """'Social Media Activity – Kick Subs - October | S1-S402' -> 'Kick Subs - October'."""
    m = re.match(r"^.*?–\s*(.*?)\s*\|", text(ws, ref))
    return m.group(1) if m else ""


def prefix_of(ws, ref):
    return text(ws, ref).split("–")[0].strip()  # "Social Media Activity" / "Projects Activity"


def slots(ws, c):
    """Linhas de itens do bloco que começa na coluna c (nome mesclado em 4 colunas x 2 linhas).
    As linhas vazias não são ocultadas: nos torneios os blocos dividem as linhas e, nas outras abas,
    a caixa da justificativa ocupa as mesmas linhas dos últimos itens."""
    rows = []
    merged = {str(m) for m in ws.merged_cells.ranges}
    r = 15
    while f"{c}{r}:{col(c, 3)}{r + 1}" in merged:
        rows.append(r)
        r += 2
    return rows


def item_group(ws, c, hide):
    out = []
    for r in slots(ws, c):
        s = {"cells": [
            {"cell": f"{c}{r}", "format": "{label}", "rich": "itemLabel"},
            {"cell": f"{col(c, 5)}{r}", "format": "{id}", "rich": "plain"},
            {"cell": f"{col(c, 6)}{r}", "format": "*{qty}", "rich": "value"},
        ], "image": f"{col(c, 4)}{r}:{col(c, 4)}{r + 1}"}
        if hide:
            s["hideRows"] = [r, r + 1]
        out.append(s)
    return {"key": "items", "label": "Itens", "kind": "items", "labelStyle": "twoLines", "slots": out}


def stamp_cell(ws):
    for row in ws.iter_rows(min_row=4, max_row=8):
        for c in row:
            if isinstance(c.value, str) and re.fullmatch(r"\s*s1\s*-\s*s?\d+\s*", c.value):
                return c.coordinate
    return None


def block_fields(ws, c, prefix, dates_from_section):
    fl = [
        field("activity", "Nome da atividade", f"{c}4", "text", f"{prefix} – {{value}} | {{serversUpper}}", default=activity_name(ws, f"{c}4")),
        field("quantity", "Quantidade de códigos", f"{c}9", "text", 'QUANTITY OF "{field.activity}" CODES: [{value}].'),
        field("tabName", "Nome da aba (no jogo)", f"{col(c, 2)}12", "text", default=text(ws, f"{col(c, 2)}12")),
    ]
    if dates_from_section:
        fl += [
            field("startCopy", "Início", f"{c}6", "derived", "DATE ENTRY: ({start.date}) - {start.time}", "date"),
            field("endCopy", "Fim", f"{c}7", "derived", "DATE END: ({end.date}) - {end.time}", "date"),
        ]
    return fl


layouts = []


def layout(sheet, lid, type_, label, description, pattern, blocks_cols, block_names=None):
    ws = wb[sheet]
    prefix = prefix_of(ws, "D4")
    fields = [
        field("start", "Início", "D6", "datetime", "DATE ENTRY: ({date}) - {time}", "date"),
        field("end", "Fim", "D7", "datetime", "DATE END: ({date}) - {time}", "date"),
    ]
    stamp = stamp_cell(ws)
    if stamp:
        fields.append(field("serversStamp", "Servidores (selo)", stamp, "derived", "{servers}"))
    single = len(blocks_cols) == 1
    blocks = []
    for i, c in enumerate(blocks_cols):
        blocks.append({
            "fields": block_fields(ws, c, prefix, dates_from_section=i > 0),
            "groups": [item_group(ws, c, hide=False)],
        })
    d = {
        "id": lid, "type": type_, "label": label, "description": description,
        "sheet": sheet, "sheetNamePattern": pattern,
        "fields": fields,
        "sets": [{"key": "prizes", "label": "Premiações" if not single else "Códigos", "blockLabel": "Premiação" if not single else "Código",
                  "minBlocks": len(blocks), "blocks": blocks}],
        "signature": sorted(str(m) for m in ws.merged_cells.ranges if m.min_row <= 12)[:40],
    }
    if block_names:
        d["blockNames"] = block_names
    layouts.append(d)


layout("Solicitação Código KickSub", "code-kicksub", "code", "Solicitação Código KickSub",
       "Códigos para os subs da Kick: datas, quantidade, nome da aba e até 9 itens.", "Solicitação Código KickSub", ["D"])
layout("Solicitação Código Lives", "code-lives", "code", "Solicitação Código Lives",
       "Códigos das lives: datas, quantidade, nome da aba e até 8 itens.", "Solicitação Código Lives", ["D"])
layout("Solicitação Código 3", "code-tournament-5", "code",
       "Solicitação Código Torneios (Premiação para top1, 2, 3, 4-16 e participação)",
       "Competição com 5 premiações: 1º, 2º, 3º, 4º ao 16º lugar e participação.", "Código Torneio Top1-3,4-16,Part",
       ["D", "L", "T", "AB", "AJ"], ["1º lugar", "2º lugar", "3º lugar", "4º ao 16º lugar", "Participação"])
layout("Solicitação Código 4", "code-tournament-4", "code",
       "Solicitação Código Torneios (Premiação para top1, 2, 3 e participação)",
       "Competição com 4 premiações: 1º, 2º, 3º lugar e participação.", "Código Torneio Top1-3 e Part",
       ["D", "L", "T", "AB"], ["1º lugar", "2º lugar", "3º lugar", "Participação"])

digest = hashlib.sha256(open(SRC, "rb").read()).hexdigest()
with open(OUT, "w", encoding="utf-8") as f:
    json.dump({"version": VERSION, "sha256": digest, "fileName": FILE_NAME, "layouts": layouts}, f, ensure_ascii=False, separators=(",", ":"))
print(f"{len(layouts)} layouts -> {OUT} ({digest[:12]})")
